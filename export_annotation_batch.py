"""
Export a batch of real articles into two independent annotation spreadsheets
— one for each annotator — for building the 200+ item gold-standard
benchmark (BRD Stage 4 / REQ-8's "Benchmark Information").

Why two separate FILES rather than two tabs in one workbook: Cohen's kappa
measures agreement between annotators who labeled *independently*. Two tabs
in the same file make it too easy to glance at the other column while
working. Two files, sent separately, keep the labeling genuinely blind.

Usage:
    python export_annotation_batch.py
    python export_annotation_batch.py --count 250
"""

import argparse
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.db import get_session
from app.models import Article, Entity

# The 5 headlines bundled in app/ingestion/sample_data.py, used as a
# ingestion-testing fallback when no NEWSAPI_KEY is set. These are excluded
# from the benchmark since scoring FinBERT against deliberately clear-cut
# canned headlines wouldn't be a meaningful academic evaluation — the
# benchmark needs real, sometimes-ambiguous live articles.
_SAMPLE_HEADLINES = {
    "Apple beats quarterly earnings expectations on strong iPhone demand",
    "Tesla shares slide after production miss and price cut announcement",
    "Microsoft unveils new AI features across Office suite, stock hits record high",
    "NVIDIA faces new export restrictions on advanced AI chips to China",
    "Amazon announces $10 billion investment in logistics automation",
}

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
EDITABLE_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
BODY_FONT = Font(name="Arial")


def fetch_candidate_articles(count: int) -> list[dict]:
    with get_session() as session:
        articles = (
            session.query(Article)
            .filter(Article.body.isnot(None))
            .filter(Article.headline.notin_(_SAMPLE_HEADLINES))
            .all()
        )

    rows = [
        {
            "article_id": a.article_id,
            "headline": a.headline,
            "body_preview": (a.body[:1000] + "…") if a.body and len(a.body) > 1000 else (a.body or ""),
            "source_name": a.source_name or "",
            "published_at": a.published_at,
        }
        for a in articles
    ]

    random.Random(42).shuffle(rows)  # fixed seed: reproducible, not chronologically biased
    return rows[:count]


def fetch_entity_reference() -> list[str]:
    with get_session() as session:
        return [name for (name,) in session.query(Entity.name).filter(Entity.is_active.is_(True)).all()]


def build_instructions_sheet(wb: Workbook, annotator_label: str, entity_names: list[str], total_items: int) -> None:
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 100

    lines = [
        f"FinSight AI — Gold-Standard Annotation ({annotator_label})",
        "",
        f"You have {total_items} articles to label in the 'Annotations' tab.",
        "Work independently — do not compare answers with the other annotator until both of you are done.",
        "",
        "For each article, fill in two columns:",
        "  1. Sentiment — choose exactly one: positive, negative, or neutral",
        "     (a dropdown is provided in that column)",
        "  2. Entities Mentioned — type the company/index/sector/commodity names",
        "     actually mentioned in the article, separated by commas.",
        "     Use the exact names below so they match later.",
        "",
        "How to judge sentiment:",
        "  positive = the news is favorable for the company/entity's outlook",
        "  negative = the news is unfavorable for the company/entity's outlook",
        "  neutral  = factual/mixed with no clear favorable or unfavorable lean",
        "",
        "Example (do not edit — for reference only):",
        "  Headline: 'Apple reports record iPhone sales, beats forecasts'",
        "  Sentiment: positive",
        "  Entities Mentioned: Apple Inc.",
        "",
        "Known entity names (for consistent spelling):",
        "  " + ", ".join(entity_names),
    ]
    for i, line in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=line)
        cell.font = Font(name="Arial", bold=(i == 1), size=13 if i == 1 else 11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_annotation_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("Annotations")
    headers = ["Article ID", "Headline", "Body Preview", "Source", "Published At",
               "Sentiment (positive/negative/neutral)", "Entities Mentioned"]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for r, row in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=row["article_id"]).font = BODY_FONT
        ws.cell(row=r, column=2, value=row["headline"]).font = BODY_FONT
        body_cell = ws.cell(row=r, column=3, value=row["body_preview"])
        body_cell.font = BODY_FONT
        body_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=4, value=row["source_name"]).font = BODY_FONT
        published = row["published_at"]
        ws.cell(row=r, column=5, value=published.strftime("%Y-%m-%d %H:%M") if published else "").font = BODY_FONT

        for col in (6, 7):
            cell = ws.cell(row=r, column=col, value="")
            cell.font = BODY_FONT
            cell.fill = EDITABLE_FILL

    dv = DataValidation(type="list", formula1='"positive,negative,neutral"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Invalid entry",
                        error="Choose positive, negative, or neutral.")
    ws.add_data_validation(dv)
    dv.add(f"F2:F{len(rows) + 1}")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 22
    ws.column_dimensions["G"].width = 30
    ws.freeze_panes = "A2"


def build_workbook(annotator_label: str, rows: list[dict], entity_names: list[str]) -> Workbook:
    wb = Workbook()
    build_instructions_sheet(wb, annotator_label, entity_names, len(rows))
    build_annotation_sheet(wb, rows)
    return wb


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export the gold-standard annotation batch.")
    parser.add_argument("--count", type=int, default=200)
    args = parser.parse_args()

    rows = fetch_candidate_articles(args.count)
    entity_names = fetch_entity_reference()

    if len(rows) < 200:
        print(f"WARNING: only {len(rows)} eligible articles found (BRD requires 200+). "
              f"Run more ingestion cycles before finalizing the benchmark.")
    else:
        print(f"Exporting {len(rows)} articles.")

    for label, filename in [("Annotator 1", "annotation_batch_annotator1.xlsx"),
                            ("Annotator 2", "annotation_batch_annotator2.xlsx")]:
        wb = build_workbook(label, rows, entity_names)
        wb.save(filename)
        print(f"Saved {filename}")

    print("\nSend one file to each annotator. Do not let them compare answers until both are complete.")
